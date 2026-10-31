import io
import json
import zipfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.services.csv_loader import CSVDataLoader
from app.services.dataset_readers import DatasetReadError, detect_format, ingest_dataset, preview_dataset
from app.services.mobility_store import MobilityStore
from app.services.dataset_insights import overview, geographic_features
from tests.test_ingestion import gtfs_archive


def ingest(path, tmp_path, metadata=None):
    store=MobilityStore(tmp_path/'store')
    result=ingest_dataset(path,path.name,store,metadata or {})
    return store,result


@pytest.mark.parametrize('suffix,content',[
    ('.csv','LAT,Lng,Recorded At,Device ID\n17.4,78.4,2025-02-01T08:30:00,device1\n'),
    ('.tsv','LAT\tLng\tRecorded At\n17.4\t78.4\t2025-02-01T08:30:00\n'),
    ('.txt','LAT;Lng;Recorded At\n17.4;78.4;2025-02-01T08:30:00\n'),
    ('.json',json.dumps([{'lat':17.4,'lng':78.4,'recorded_at':'2025-02-01T08:30:00'}])),
    ('.jsonl','{"lat":17.4,"lng":78.4,"recorded_at":"2025-02-01T08:30:00"}\n{"lat":17.5,"lng":78.5,"recorded_at":"2025-02-01T09:30:00"}\n'),
    ('.ndjson','{"lat":17.4,"lng":78.4,"recorded_at":"2025-02-01T08:30:00"}\n'),
])
def test_tabular_formats_reach_catalog_and_analyses(tmp_path,suffix,content):
    path=tmp_path/f'fixture{suffix}';path.write_text(content,encoding='utf-8')
    preview=preview_dataset(path)
    assert preview['requires_mapping'] is False
    store,result=ingest(path,tmp_path)
    assert result['record_count']>=1
    assert store.spatial_cells(path.name)['total_records']>=1
    assert store.temporal_summary(path.name)['hourly']


def test_excel_readers(tmp_path):
    path=tmp_path/'fixture.xlsx';book=openpyxl.Workbook();sheet=book.active
    sheet.append(['lat','lon','recorded_at']);sheet.append([17.4,78.4,'2025-01-01T08:00:00']);book.save(path)
    store,result=ingest(path,tmp_path)
    assert result['format']=='xlsx' and store.temporal_summary(path.name)['total_records']==1
    legacy=Path(__file__).parent/'fixtures'/'tabular.xls'
    store,result=ingest(legacy,tmp_path)
    assert result['format']=='xls' and result['record_count']==1
    assert store.spatial_cells(legacy.name)['total_records']==1
    renamed=tmp_path/'download.csv';renamed.write_bytes(path.read_bytes())
    store,result=ingest(renamed,tmp_path)
    assert result['format']=='xlsx' and result['record_count']==1


def test_single_sheet_requirement_is_explicit(tmp_path):
    path=tmp_path/'multiple.xlsx';book=openpyxl.Workbook()
    for sheet in [book.active,book.create_sheet('other')]:sheet.append(['city','value']);sheet.append(['Hyderabad',1])
    book.save(path)
    with pytest.raises(DatasetReadError,match='one non-empty data sheet'):preview_dataset(path)


def test_static_coordinates_do_not_require_or_create_timestamps(tmp_path):
    path=tmp_path/'stops.csv';path.write_text('lat,lon,name\n17.4,78.4,Park\n')
    store,result=ingest(path,tmp_path)
    assert 'spatial' in result['supported_analyses'] and 'temporal' not in result['supported_analyses']
    assert store.page_observations(path.name)[0][0]['timestamp'] is None
    assert store.temporal_summary(path.name)['status']=='no_data'


def test_urban_indicators_have_no_fabricated_gps_and_keep_units_separate(tmp_path):
    path=tmp_path/'cities.csv';path.write_text('City,Indicator,Value,Year,Unit\nHyderabad,roads,12,2025,km\nHyderabad,roads,7,2025,count\nPune,roads,15,2025,km\n')
    store,result=ingest(path,tmp_path)
    summary=overview(store,path.name)
    assert result['record_count']==3 and 'indicators' in result['supported_analyses']
    assert len(summary['indicators'])==3 and summary['totals']['geocoded_records']==0
    assert all(row['timestamp'] is None for row in store.page_observations(path.name)[0])


def test_timestamp_only_dataset_and_date_only_precision(tmp_path):
    path=tmp_path/'events.csv';path.write_text('date,count\n2025-01-01,4\n2025-01-02,2\n')
    store,result=ingest(path,tmp_path)
    temporal=store.temporal_summary(path.name)
    assert temporal['total_records']==2 and temporal['hourly']=={}
    assert sum(temporal['daily'].values())==6 and 'spatial' not in result['supported_analyses']


def test_clock_time_without_date_is_rejected(tmp_path):
    path=tmp_path/'clock.csv';path.write_text('lat,lon,time\n17.4,78.4,08:30:00\n')
    with pytest.raises(DatasetReadError,match='No valid mobility rows'):ingest(path,tmp_path)


def test_ambiguous_mapping_requires_confirmation(tmp_path):
    path=tmp_path/'ambiguous.csv';path.write_text('latitude,lat,lon\n17.4,17.5,78.4\n')
    preview=preview_dataset(path)
    assert 'latitude' in preview['ambiguous_fields']
    with pytest.raises(DatasetReadError,match='Confirm schema'):ingest(path,tmp_path)
    store,_=ingest(path,tmp_path,{'column_mapping':{'latitude':'lat','longitude':'lon'}})
    assert store.page_observations(path.name)[0][0]['latitude']==17.5


def test_xy_confirmation_and_profile_persistence(tmp_path):
    path=tmp_path/'xy.tsv';path.write_text('x\ty\n78.4\t17.4\n')
    assert preview_dataset(path)['requires_mapping']
    store,_=ingest(path,tmp_path,{'column_mapping':{'latitude':'y','longitude':'x'}})
    store.save_mapping_profile('Survey',{'latitude':'y','longitude':'x'},'geographic',None)
    assert store.mapping_profiles()[0]['mapping']['latitude']=='y'


def test_ambiguous_date_order_is_not_silently_assumed(tmp_path):
    path=tmp_path/'date.csv';path.write_text('lat,lon,datetime\n17.4,78.4,03/04/2025 08:30\n')
    with pytest.raises(DatasetReadError):ingest(path,tmp_path)
    store,_=ingest(path,tmp_path,{'date_format':'%d/%m/%Y %H:%M'})
    assert store.page_observations(path.name)[0][0]['local_day']=='2025-04-03'


def test_nested_generic_zip_and_content_detection(tmp_path):
    path=tmp_path/'tables.zip'
    with zipfile.ZipFile(path,'w') as archive:archive.writestr('nested/data.tsv','lat\tlon\n17.4\t78.4\n')
    store,result=ingest(path,tmp_path)
    assert result['format']=='tabular_zip' and result['record_count']==1
    renamed=tmp_path/'renamed.csv';renamed.write_bytes(gtfs_archive(tmp_path/'feed.zip',prefix='nested/feed/'))
    store,result=ingest(renamed,tmp_path)
    assert result['format']=='gtfs_static'


def test_archives_with_unrelated_schemas_are_not_silently_joined(tmp_path):
    path=tmp_path/'mixed.zip'
    with zipfile.ZipFile(path,'w') as archive:
        archive.writestr('a.csv','lat,lon\n17.4,78.4\n')
        archive.writestr('b.csv','city,indicator,value\nHyderabad,parks,5\n')
    with pytest.raises(DatasetReadError,match='different table schemas'):ingest(path,tmp_path)
    assert preview_dataset(path)['requires_table_selection']
    store,result=ingest(path,tmp_path/'selected',{'archive_member':'b.csv'})
    assert result['record_count']==1 and overview(store,path.name)['indicators'][0]['indicator']=='parks'


@pytest.mark.parametrize('case',['empty','traversal','symlink','corrupt','duplicate'])
def test_unsafe_archives(case,tmp_path):
    path=tmp_path/'bad.zip'
    if case=='corrupt':path.write_bytes(b'PK broken')
    else:
        with zipfile.ZipFile(path,'w') as archive:
            if case=='traversal':archive.writestr('a/../../table.csv','lat,lon\n17,78\n')
            if case=='symlink':
                info=zipfile.ZipInfo('link.csv');info.external_attr=0o120777<<16;archive.writestr(info,'target')
            if case=='duplicate':
                archive.writestr('a.csv','lat,lon\n17,78\n');archive.writestr('a.csv','lat,lon\n17,78\n')
    with pytest.raises(DatasetReadError):detect_format(path)


def test_geometry_only_geojson_becomes_active_and_available(tmp_path):
    path=tmp_path/'boundary.geojson';path.write_text(json.dumps({'type':'Feature','geometry':{'type':'Polygon','coordinates':[[[78.4,17.4],[78.5,17.4],[78.5,17.5],[78.4,17.4]]]},'properties':{'city':'Hyderabad'}}))
    loader=CSVDataLoader(str(tmp_path/'catalog'))
    result=loader.import_dataset('boundary.geojson',path.read_bytes())
    assert result['active'] and 'geometry' in result['supported_analyses']
    assert geographic_features(loader.store,'boundary.geojson')['features'][0]['geometry']['type']=='Polygon'


def test_gtfs_optional_tables_relationships_frequency_and_report(tmp_path):
    path=tmp_path/'service.zip'
    gtfs_archive(path,prefix='feed/',override={
        'agency.txt':'agency_name,agency_url,agency_timezone\nTransit,https://example.org,Asia/Kolkata\n',
        'frequencies.txt':'trip_id,start_time,end_time,headway_secs\nt1,08:00:00,10:00:00,600\n',
        'shapes.txt':'shape_id,shape_pt_lat,shape_pt_lon,shape_pt_sequence\na,17.4,78.4,1\na,17.5,78.5,2\n',
        'feed_info.txt':'feed_publisher_name,feed_publisher_url,feed_lang\nTransit,https://example.org,en\n'})
    store,result=ingest(path,tmp_path)
    report=result['import_report']['gtfs']
    assert report['row_counts']['shapes.txt']==2 and not report['errors']
    assert overview(store,path.name)['frequency']['minimum_headway_seconds']==600
    assert geographic_features(store,path.name)['features'][0]['geometry']['type']=='LineString'
    with store.connect() as db:assert db.execute('SELECT COUNT(*) FROM gtfs_trips WHERE source_id=?',(path.name,)).fetchone()[0]==1


@pytest.mark.parametrize('override',[
    {'frequencies.txt':'trip_id,start_time,end_time,headway_secs\nmissing,08:00:00,10:00:00,600\n'},
    {'frequencies.txt':'trip_id,start_time,end_time,headway_secs\nt1,08:00:00,10:00:00,0\n'},
    {'calendar_dates.txt':'service_id,date,exception_type\nsvc,20250230,1\n'},
    {'calendar_dates.txt':'service_id,date,exception_type\nsvc,20250201,3\n'},
])
def test_invalid_gtfs_optional_tables(tmp_path,override):
    path=tmp_path/'invalid.zip';gtfs_archive(path,override=override)
    with pytest.raises(DatasetReadError):ingest(path,tmp_path)


def test_preview_import_catalog_and_report_api_workflows(tmp_path,monkeypatch):
    from app.main import app
    from app.auth import current_user, User
    monkeypatch.setitem(app.dependency_overrides, current_user, lambda: User('test-admin','test@example.invalid','admin','fixture'))
    from app.routers import data,analytics
    loader=CSVDataLoader(str(tmp_path/'catalog'))
    monkeypatch.setattr(data,'data_loader',loader);monkeypatch.setattr(analytics,'data_loader',loader)
    client=TestClient(app)
    content=b'origin,destination,count\nHyderabad,Pune,12\nHyderabad,Pune,6\n'
    preview=client.post('/api/data/preview?filename=flows.tsv',content=content)
    assert preview.status_code==200 and preview.json()['preview']['dataset_type']=='origin_destination'
    response=client.post('/api/data/import?filename=flows.tsv',content=content)
    assert response.status_code==200
    assert client.get('/api/data').json()['data']['metadata']['record_count']==2
    corridor=client.post('/api/analytics/corridors',json={}).json()['result']
    assert corridor['top_corridors'][0]['trip_count']==18
    report=client.post('/api/analytics/reports',json={})
    assert report.status_code==200 and report.json()['report']['overview']['totals']['records']==2
    assert len(client.get('/api/analytics/reports').json()['reports'])==1
    assert client.delete('/api/analytics/reports',params={'report_id':report.json()['report']['id']}).status_code==200
    assert not client.get('/api/analytics/reports').json()['reports']


def test_generic_stops_filename_is_not_mistaken_for_gtfs(tmp_path):
    path=tmp_path/'generic.zip'
    with zipfile.ZipFile(path,'w') as archive:archive.writestr('nested/stops.txt','lat,lon\n17.4,78.4\n')
    store,result=ingest(path,tmp_path)
    assert result['format']=='tabular_zip' and result['record_count']==1


def test_geographic_property_mapping_preview(tmp_path):
    path=tmp_path/'mapped.geojson'
    path.write_text(json.dumps({'type':'Feature','properties':{'capture_moment':'2025-01-01T08:00:00'},'geometry':{'type':'Point','coordinates':[78.4,17.4]}}))
    preview=preview_dataset(path)
    assert 'capture_moment' in [column['name'] for column in preview['columns']]
    mapping={**preview['mapping'],'timestamp':'capture_moment'}
    store,_=ingest(path,tmp_path,{'column_mapping':mapping})
    assert store.temporal_summary(path.name)['hourly']=={8:1}


@pytest.mark.parametrize('timestamp',['2025','2025-01',1700000000])
def test_partial_or_numeric_dates_are_not_fabricated(tmp_path,timestamp):
    path=tmp_path/'partial.json'
    path.write_text(json.dumps([{'lat':17.4,'lon':78.4,'timestamp':timestamp}]))
    with pytest.raises(DatasetReadError,match='No valid mobility rows'):ingest(path,tmp_path)


def test_compression_bomb_and_member_limits(tmp_path,monkeypatch):
    from app.services import dataset_readers
    path=tmp_path/'bomb.zip'
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:archive.writestr('table.csv',b'a'*(11*1024*1024))
    with pytest.raises(DatasetReadError,match='compression ratio'):detect_format(path)
    path=tmp_path/'member.zip'
    with zipfile.ZipFile(path,'w') as archive:archive.writestr('table.csv','lat,lon\n17.4,78.4\n')
    monkeypatch.setattr(dataset_readers,'MAX_ZIP_MEMBER_BYTES',5)
    with pytest.raises(DatasetReadError,match='uncompressed limit'):detect_format(path)


def test_catalog_activation_and_failure_preserve_previous_source(tmp_path,monkeypatch):
    from app.main import app
    from app.auth import current_user, User
    monkeypatch.setitem(app.dependency_overrides, current_user, lambda: User('test-admin','test@example.invalid','admin','fixture'))
    from app.routers import data
    loader=CSVDataLoader(str(tmp_path/'catalog'))
    loader.import_dataset('first.csv',b'lat,lon\n17.4,78.4\n')
    loader.import_dataset('second.csv',b'lat,lon\n18.4,79.4\n')
    monkeypatch.setattr(data,'data_loader',loader)
    response=TestClient(app).post('/api/data/activate',params={'filename':'first.csv'})
    assert response.status_code==200 and loader.get_active_dataset()=='first.csv'
    with pytest.raises(DatasetReadError):loader.import_dataset('bad.zip',b'PK corrupt')
    assert loader.get_active_dataset()=='first.csv' and loader.store.get_source('bad.zip') is None


def test_gtfs_agency_reference_and_calendar_range_are_validated(tmp_path):
    path=tmp_path/'agency.zip'
    gtfs_archive(path,override={'agency.txt':'agency_id,agency_name,agency_url,agency_timezone\na1,Transit,https://example.org,Asia/Kolkata\n',
                              'routes.txt':'route_id,route_type,agency_id\nr1,3,unknown\n'})
    with pytest.raises(DatasetReadError,match='agency_id'):ingest(path,tmp_path)


def test_connection_context_closes_before_temp_directory_removal(tmp_path):
    store=MobilityStore(tmp_path/'catalog')
    with store.connect() as connection:connection.execute('SELECT 1')
    import sqlite3
    with pytest.raises(sqlite3.ProgrammingError):connection.execute('SELECT 1')


def test_static_line_is_geometry_not_an_observed_journey(tmp_path):
    path=tmp_path/'road.geojson'
    path.write_text(json.dumps({'type':'Feature','geometry':{'type':'LineString','coordinates':[[78.4,17.4],[78.5,17.5]]},'properties':{'timestamp':'2025-01-01T08:00:00'}}))
    store,result=ingest(path,tmp_path)
    assert result['record_count']==1 and 'route_geometry' in result['supported_analyses']
    assert store.page_observations(path.name)[0]==[]
    assert store.temporal_summary(path.name)['total_records']==0
    assert store.od_summary(path.name)['status']=='insufficient_data'


def test_geojson_per_point_times_support_identified_trajectory(tmp_path):
    path=tmp_path/'track.geojson'
    path.write_text(json.dumps({'type':'Feature','geometry':{'type':'LineString','coordinates':[[78.4,17.4],[78.5,17.5]]},'properties':{'timestamps':['2025-01-01T08:00:00','2025-01-01T08:30:00']}}))
    store,result=ingest(path,tmp_path)
    assert result['record_count']==2
    assert store.od_summary(path.name)['data_semantics']=='trajectory_endpoint_inference'


def test_static_gtfs_summary_cache_invalidates_when_source_changes(tmp_path):
    path=tmp_path/'cached.zip';gtfs_archive(path)
    store,_=ingest(path,tmp_path)
    initial=store.temporal_summary(path.name)
    assert store.temporal_summary(path.name)==initial
    store.insert_gtfs_stop_times(path.name,[('t1','s1',3,27*3600,27*3600)])
    store.finalize_source(path.name)
    updated=store.temporal_summary(path.name)
    assert updated['total_records']==initial['total_records']+1
    assert updated['hourly']['3']==1
    assert store.delete_source(path.name)
    with store.connect() as db:assert db.execute('SELECT COUNT(*) FROM analysis_cache').fetchone()[0]==0


@pytest.mark.parametrize('member',['folder/NUL.csv','folder/data.csv:stream'])
def test_windows_archive_device_names_and_stream_paths_are_rejected(tmp_path,member):
    path=tmp_path/'windows.zip'
    with zipfile.ZipFile(path,'w') as archive:archive.writestr(member,'lat,lon\n17.4,78.4\n')
    with pytest.raises(DatasetReadError,match='unsafe path'):detect_format(path)
