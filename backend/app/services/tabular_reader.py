"""Bounded, content-aware readers for delimited, Excel, JSON and JSON Lines data."""
from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path
from typing import Iterator

import pandas as pd

from app.services.dataset_errors import DatasetReadError

MAX_DOCUMENT_BYTES = 256 * 1024 * 1024


def detect_tabular(path: Path) -> str:
    with path.open('rb') as stream: sample = stream.read(32768)
    if not sample: raise DatasetReadError('Could not parse CSV or tabular content: the uploaded file is empty.')
    if sample.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'): return 'xls'
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            if '[Content_Types].xml' in archive.namelist() and any(name.startswith('xl/') for name in archive.namelist()): return 'xlsx'
        return 'archive'
    if path.suffix.lower() in ('.zip', '.xlsx', '.xls'):
        raise DatasetReadError('The uploaded file is not a valid, readable ZIP/Excel archive.')
    try: text = sample.decode('utf-8-sig')
    except UnicodeDecodeError as exc: raise DatasetReadError('Tabular text must use UTF-8 encoding; binary content was not recognised.') from exc
    stripped = text.lstrip()
    if stripped.startswith(('{', '[')):
        if path.suffix.lower() in ('.jsonl', '.ndjson'): return 'jsonl'
        try:
            first = json.loads(stripped.splitlines()[0])
            if isinstance(first, dict) and len(stripped.splitlines()) > 1 and stripped.splitlines()[1].lstrip().startswith('{'): return 'jsonl'
        except (ValueError, IndexError): pass
        if '"FeatureCollection"' in text or '"Feature"' in text or '"coordinates"' in text: return 'json'
        return 'json'
    if stripped.startswith('Geolife trajectory') or (path.suffix.lower() == '.plt'): return 'geolife_plt'
    if '\x00' in text: raise DatasetReadError('Unsupported binary content; choose a supported tabular or geographic format.')
    try: dialect = csv.Sniffer().sniff(text, delimiters=',\t;|')
    except csv.Error:
        if len(stripped.splitlines()[0].split()) >= 2: return 'whitespace'
        raise DatasetReadError('No supported delimited table or structured JSON content was detected.')
    return 'tsv' if dialect.delimiter == '\t' else 'csv'


def frames(path: Path, fmt: str, *, sample_only: bool = False) -> Iterator[pd.DataFrame]:
    try:
        if fmt in ('csv', 'tsv', 'whitespace'):
            with path.open('r', encoding='utf-8-sig', newline='') as stream:
                sample = stream.read(32768)
            delimiter = r'\s+' if fmt == 'whitespace' else csv.Sniffer().sniff(sample, delimiters=',\t;|').delimiter
            # Validate headers before pandas can silently rename duplicate columns.
            headers = sample.splitlines()[0].split() if fmt == 'whitespace' else next(csv.reader(io.StringIO(sample), delimiter=delimiter))
            from app.services.schema_mapping import key
            if len({key(name) for name in headers}) != len(headers): raise DatasetReadError('Duplicate column names are not allowed.')
            reader = pd.read_csv(path, sep=delimiter, encoding='utf-8-sig', chunksize=25000,
                                 dtype=object, keep_default_na=False, on_bad_lines='error')
            for frame in reader:
                yield frame.head(200) if sample_only else frame
                if sample_only: break
        elif fmt in ('xlsx', 'xls'):
            if path.stat().st_size > MAX_DOCUMENT_BYTES: raise DatasetReadError('Excel workbooks must be 256 MB or smaller.')
            if fmt == 'xlsx':
                import openpyxl
                binary = path.open('rb')
                try: workbook = openpyxl.load_workbook(binary, read_only=True, data_only=True)
                except Exception:
                    binary.close()
                    raise
                try:
                    # Multiple non-empty sheets need a deliberate choice, not a silent merge.
                    sheets = [sheet for sheet in workbook if sheet.max_row and sheet.max_row > 1]
                    if len(sheets) != 1: raise DatasetReadError('Excel imports require one non-empty data sheet. Export the intended sheet or provide a single-table workbook.')
                    rows = sheets[0].iter_rows(values_only=True); headers = next(rows, None)
                    if not headers: raise DatasetReadError('Excel sheet has no header row.')
                    batch = []
                    for row in rows:
                        if not any(value is not None for value in row): continue
                        batch.append(row)
                        if len(batch) >= (200 if sample_only else 25000):
                            yield pd.DataFrame(batch, columns=[str(h or '') for h in headers]); batch=[]
                            if sample_only: return
                    if batch: yield pd.DataFrame(batch, columns=[str(h or '') for h in headers])
                finally:
                    workbook.close()
                    binary.close()
            else:
                workbook = pd.ExcelFile(path, engine='xlrd')
                try:
                    if len(workbook.sheet_names) != 1: raise DatasetReadError('XLS imports require one data sheet; export the intended sheet first.')
                    yield pd.read_excel(workbook, dtype=object, keep_default_na=False, nrows=200 if sample_only else None)
                finally: workbook.close()
        elif fmt == 'jsonl':
            batch=[]
            with path.open('r', encoding='utf-8-sig') as stream:
                for index, line in enumerate(stream, 1):
                    if not line.strip(): continue
                    if len(line) > 8 * 1024 * 1024: raise DatasetReadError(f'JSON Lines row {index} exceeds 8 MB.')
                    try: row=json.loads(line)
                    except ValueError as exc: raise DatasetReadError(f'Invalid JSON Lines at row {index}: {exc}') from exc
                    if not isinstance(row, dict): raise DatasetReadError(f'JSON Lines row {index} must be an object.')
                    batch.append(row)
                    if len(batch)>=(200 if sample_only else 25000):
                        yield pd.DataFrame(batch); batch=[]
                        if sample_only:return
            if batch:yield pd.DataFrame(batch)
        elif fmt == 'json':
            if path.stat().st_size > MAX_DOCUMENT_BYTES: raise DatasetReadError('JSON documents must be 256 MB or smaller; use JSON Lines for larger tables.')
            document=json.loads(path.read_text(encoding='utf-8-sig'))
            if isinstance(document, dict):
                document = document.get('records', document.get('data', document))
                if isinstance(document, dict): document=[document]
            if not isinstance(document,list) or not all(isinstance(row,dict) for row in document):
                raise DatasetReadError('Tabular JSON must contain objects, an array of objects, or a records/data array.')
            for start in range(0,len(document),25000):
                yield pd.DataFrame(document[start:start+(200 if sample_only else 25000)])
                if sample_only:break
        else: raise DatasetReadError(f'No tabular reader for {fmt}.')
    except DatasetReadError: raise
    except Exception as exc: raise DatasetReadError(f'Could not parse {fmt.upper()} {path.name}: {exc}') from exc


def clean_rows(frame: pd.DataFrame) -> list[dict]:
    # Preserve nested JSON properties while replacing pandas null sentinels.
    return json.loads(frame.to_json(orient='records', date_format='iso'))
